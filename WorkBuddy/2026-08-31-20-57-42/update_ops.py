from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

SRC="/Users/temu/WorkBuddy/2026-08-31-20-57-42/733汇报_8月总结与行动计划.xlsx"
wb=load_workbook(SRC)

LEFT=Alignment(horizontal="left",vertical="center",wrap_text=True)
TITLE=Font(bold=True,size=14,name="微软雅黑",color="1F3864")

def rebuild(sheet_name, title, items):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws=wb.create_sheet(sheet_name)
    ws["A1"]=title; ws["A1"].font=TITLE
    r=3
    for it in items:
        c=ws.cell(row=r,column=1,value=it); c.alignment=LEFT
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=1)
        lines=it.count("\n")+1
        ws.row_dimensions[r].height=max(34, lines*16)
        r+=1
    ws.column_dimensions["A"].width=118
    return ws

month_ops=[
 "1. 热销款 & 睡袍维护（重点动作）\n"
 "• 7 款热销链接重点维护：盯限流情况、库存、流量\n"
 "• 标红 8 款 SPU 安排重新上架（库存冗余）\n"
 "• 12 款睡袍：维护好链接 + 核价 + 找对应买手跟踪链接情况，关注 9 月销售\n"
 "  上架安排不限于：①单件装（所有睡袍款式已完成并前台在售，单色单款出单即复色） ②双件装 ③情侣款裂变\n"
 "  核价安排：找买手核价争取目标核价（目前男士/女士睡袍买手已对接好）",

 "2. 入仓新品推新 & 动销观察（清单 8 款 SKU）\n"
 "• 关注本月即将入仓产品数量，维护热销库存、加大活动销售\n"
 "• 重新推新产品，观察动销：WAW8549 / WAN1196 / WAW7481 / WAN1259 / WAW8649 / WAW8580 / WAW7398 / WAW7480\n"
 "• 动作：①重新做图测试不同图片效果 ②找买手查看链接质量 ③查看竞品链接和图片，对标主体链接做差异化",

 "3. 新上架链接：提交广告 + 每周链接活动\n"
 "• 所有新上架链接同步提交广告计划，并按周提报链接活动",

 "4. 内裤链接限流处理（3 款）\n"
 "• 找买手对接 WAN1049 / WAN1196 / WAN1213 共 3 款内裤链接\n"
 "• 平台要求图片完全非人模展示，目前这 3 款限流严重，需找买手沟通并同步关注进度\n"
 "（注：WAN1196 同时在点 2 推新清单中，需一并跟踪）",

 "5. 链接优化策略（站内清理 + 差异化定价）\n"
 "• 站内同款互比链接已下架（本周已完成）\n"
 "• 竞品对标与差异化定价：通过竞品链接比价、核价，采用『阶梯定价 / 限定价』方式，打造差异化避免价格战内卷\n"
 "• 模仿优秀竞品玩法，打造与竞品的差异化优势，避免陷入同质化竞争",

 "6. 欧区上新 & 在售动销监控\n"
 "• 持续欧区链接上新情况\n"
 "• 关注在售链接动销：6014 / 7191 / 6798",
]

week_plan=[
 "1. 维护 7 款热销链接（限流/库存/流量）；标红 8 款 SPU 重新上架。",
 "2. 跟踪入仓入库单，维护热销库存、加大活动；推新观察 8 款（WAW8549/WAN1196/WAW7481/WAN1259/WAW8649/WAW8580/WAW7398/WAW7480），做图测效+查链接质量+竞品对标。",
 "3. 12 款睡袍：单件装已上架，推进双件装+情侣款裂变；男/女睡袍买手已对接，跟进核价目标。",
 "4. 对接买手处理 3 款内裤（WAN1049/1196/1213）非人模图片限流问题，同步进度。",
 "5. 提交新上架链接广告+每周活动；落实阶梯/限定价差异化策略（同款互比链接已下架）。",
 "6. 推进欧区上新；监控在售动销 6014/7191/6798。",
 "7. 盘点红色清仓清单，启动睡袍/死库降价清仓；确认在途 9,687 件拦截方案。",
]

rebuild("本月运营操作","本月（9月）运营操作", month_ops)
rebuild("本周运营计划","本周运营计划", week_plan)

wb.save(SRC)
print("UPDATED:",SRC)
print("Sheets:", wb.sheetnames)
