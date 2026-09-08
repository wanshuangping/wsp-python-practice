import xml.etree.ElementTree as ET
if not hasattr(ET.ElementTree, 'getiterator'):
    ET.ElementTree.getiterator = lambda self, tag=None: self.iter(tag)

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FILES = {
    "A": "/Users/temu/Downloads/2026-09-01_bizWarehouseProductInventoryExport6409.xlsx",
    "B": "/Users/temu/Downloads/ProductInventory-1788234957397-556059.xlsx",
    "C": "/Users/temu/Downloads/产品库存(2).xls",
    "master": "/Users/temu/Downloads/导出产品-按SKU-953261488470130688.xlsx",
}

def norm(s):
    return "".join(ch for ch in str(s).lower() if ch.isalnum())

def num(x):
    try:
        if x is None or (isinstance(x, float) and pd.isna(x)):
            return 0
        s = str(x).strip().replace(",", "")
        return int(float(s)) if s not in ("", "nan", "None") else 0
    except Exception:
        return 0

def find_col(df, target):
    t = norm(target)
    hm = {norm(c): c for c in df.columns}
    if t in hm:
        return hm[t]
    for k, v in hm.items():
        if t in k or k in t:
            return v
    return None

WH_NAME = {"A": "仓库A(5775)", "B": "仓库B(WPLA)", "C": "仓库C(休斯顿)"}

def parse_wh(key):
    df = pd.read_excel(FILES[key], dtype=str)  # current files have header at row 0
    sku_c = find_col(df, "sku")
    on_c = find_col(df, "在途")
    sell_c = find_col(df, "可售")
    out = []
    for _, r in df.iterrows():
        sku = str(r.get(sku_c, "")).strip() if sku_c else ""
        if not sku:
            continue
        out.append({
            "sku": sku,
            "在途": num(r.get(on_c)) if on_c else 0,
            "可售": num(r.get(sell_c)) if sell_c else 0,
        })
    return out, (sku_c, on_c, sell_c)

def load_master():
    df = pd.read_excel(FILES["master"], sheet_name="产品", dtype=str)
    sku_c = "*SKU" if "*SKU" in df.columns else df.columns[0]
    spu_c = "SPU" if "SPU" in df.columns else None
    name_c = "品名" if "品名" in df.columns else None
    exact = {}
    spu_index = []
    for _, r in df.iterrows():
        sk = str(r.get(sku_c, "")).strip()
        sp = str(r.get(spu_c, "")).strip() if spu_c else ""
        nm = str(r.get(name_c, "")).strip() if name_c else ""
        if sk:
            exact[norm(sk)] = (sk, sp, nm)
        if sp:
            spu_index.append((norm(sp), sp, nm))
    return exact, spu_index

# ---------- aggregate ----------
agg = {}
colinfo = {}
for key in ("A", "B", "C"):
    rows, info = parse_wh(key)
    colinfo[key] = info
    for row in rows:
        sk = row["sku"]
        d = agg.setdefault(sk, {w: {"在途": 0, "可售": 0} for w in ("A", "B", "C")})
        d[key]["在途"] += row["在途"]
        d[key]["可售"] += row["可售"]

exact, spu_index = load_master()

def xref(sku):
    n = norm(sku)
    if n in exact:
        return exact[n]
    for sspu, spu, nm in spu_index:
        if sspu and sspu in n:
            return (sku, spu, nm)
    return (sku, "", "")

rows = []
for sku, d in agg.items():
    a, b, c = d["A"], d["B"], d["C"]
    on_t = a["在途"] + b["在途"] + c["在途"]
    sell_t = a["可售"] + b["可售"] + c["可售"]
    sk, spu, nm = xref(sku)
    rows.append({
        "sku": sku, "spu": spu, "name": nm,
        "A_on": a["在途"], "A_sell": a["可售"],
        "B_on": b["在途"], "B_sell": b["可售"],
        "C_on": c["在途"], "C_sell": c["可售"],
        "on_t": on_t, "sell_t": sell_t,
    })

rows.sort(key=lambda r: (-(r["on_t"] + r["sell_t"]), r["sku"]))

# ---------- write single table ----------
wb = Workbook()
ws = wb.active
ws.title = "在途可售SKU汇总"
headers = ["SKU", "SPU(产品表)", "品名(产品表)",
           f"在途-{WH_NAME['A']}", f"可售-{WH_NAME['A']}",
           f"在途-{WH_NAME['B']}", f"可售-{WH_NAME['B']}",
           f"在途-{WH_NAME['C']}", f"可售-{WH_NAME['C']}",
           "在途合计", "可售合计"]
ws.append(headers)
for r in rows:
    ws.append([r["sku"], r["spu"], r["name"],
               r["A_on"], r["A_sell"], r["B_on"], r["B_sell"],
               r["C_on"], r["C_sell"], r["on_t"], r["sell_t"]])

tot = {k: sum(r[k] for r in rows) for k in ["A_on","A_sell","B_on","B_sell","C_on","C_sell","on_t","sell_t"]}
ws.append(["合计", "", f"{len(rows)}个SKU",
           tot["A_on"], tot["A_sell"], tot["B_on"], tot["B_sell"],
           tot["C_on"], tot["C_sell"], tot["on_t"], tot["sell_t"]])

# style
hdr_fill = PatternFill("solid", fgColor="4472C4")
hdr_font = Font(bold=True, color="FFFFFF", name="微软雅黑", size=10)
tot_fill = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")
left = Alignment(horizontal="left", vertical="center")

for c in range(1, len(headers)+1):
    cell = ws.cell(row=1, column=c)
    cell.fill = hdr_fill; cell.font = hdr_font; cell.alignment = center; cell.border = border
last = ws.max_row
for c in range(1, len(headers)+1):
    cell = ws.cell(row=last, column=c)
    cell.fill = tot_fill; cell.font = Font(bold=True, name="微软雅黑"); cell.border = border
    cell.alignment = center if c >= 4 else left
for rr in range(2, ws.max_row+1):
    for c in range(4, len(headers)+1):
        ws.cell(row=rr, column=c).number_format = "#,##0"
        ws.cell(row=rr, column=c).alignment = center
    for c in (1, 2, 3):
        ws.cell(row=rr, column=c).alignment = left

widths = [32, 16, 26, 14, 14, 14, 14, 16, 16, 10, 10]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"

# ---------- 说明 sheet ----------
ws2 = wb.create_sheet("说明")
matched = sum(1 for r in rows if r["spu"])
notes = [
    "在途/可售 SKU 汇总说明（基于本次附带的 3 个仓库文件）",
    "",
    f"1. 数据源：仓库A=5775 / 仓库B=WPLA / 仓库C=TX休斯顿大牛仓。",
    f"2. 共汇总 SKU {len(rows)} 个；其中 {matched} 个通过 SKU 或 SPU 片段匹配到产品表(导出产品-按SKU)，其余 SPU/品名 留空。",
    f"3. 每行 = 一个 SKU；在途/可售 为该 SKU 在三仓的合计；A/B/C 三列给出分仓明细（0 即该仓无此 SKU）。",
    f"4. 在途合计 {tot['on_t']:,} ，可售合计 {tot['sell_t']:,} 。",
    f"   分仓：仓库A 在途 {tot['A_on']:,}/可售 {tot['A_sell']:,}；仓库B 在途 {tot['B_on']:,}/可售 {tot['B_sell']:,}；仓库C 在途 {tot['C_on']:,}/可售 {tot['C_sell']:,}。",
    "5. 重要：本次仓库C(休斯顿)导出文件仅含「在途」一列，无「可售」字段，故仓库C 可售全部记为 0。若实际有可售库存，需补带可售列的导出。",
    "6. 口径：在途=采购/调拨未入库；可售=在库可销售。其余状态(待上架/待出库/不良品等)未纳入本表。",
    "7. 排序：按 (在途+可售) 总量降序。",
]
for i, t in enumerate(notes, 1):
    ws2.cell(row=i, column=1, value=t)
ws2.column_dimensions["A"].width = 115

out = "/Users/temu/WorkBuddy/2026-08-31-20-57-42/在途可售SKU汇总.xlsx"
wb.save(out)
print("SAVED:", out)
print("SKU count:", len(rows), "matched SPU:", matched)
print("在途合计:", tot["on_t"], "可售合计:", tot["sell_t"])
print("分仓 在途 A/B/C:", tot["A_on"], tot["B_on"], tot["C_on"])
print("分仓 可售 A/B/C:", tot["A_sell"], tot["B_sell"], tot["C_sell"])
print("colinfo A/B/C:", colinfo)
