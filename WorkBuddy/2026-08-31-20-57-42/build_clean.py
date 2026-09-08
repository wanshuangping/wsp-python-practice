import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = "/Users/temu/WorkBuddy/2026-08-31-20-57-42/在途可售SKU汇总.xlsx"
OUT = SRC  # overwrite in place

df = pd.read_excel(SRC, sheet_name="在途可售SKU汇总", dtype=str)
df = df[df["SKU"] != "合计"].copy()
for c in ["在途合计","可售合计"]:
    df[c] = pd.to_numeric(df[c].str.replace(",",""), errors="coerce").fillna(0).astype(int)

before = len(df)
kept = df[~((df["在途合计"] == 0) & (df["可售合计"] == 0))].copy()
removed = before - len(kept)

# per-warehouse cols for sorting/totals
wh_cols = ["在途-仓库A(5775)","可售-仓库A(5775)","在途-仓库B(WPLA)","可售-仓库B(WPLA)",
           "在途-仓库C(休斯顿)","可售-仓库C(休斯顿)","在途合计","可售合计"]
for c in wh_cols:
    kept[c] = pd.to_numeric(kept[c].astype(str).str.replace(",", ""), errors="coerce").fillna(0).astype(int)
kept["总量"] = kept["在途合计"] + kept["可售合计"]
kept = kept.sort_values("总量", ascending=False).drop(columns="总量")

# ---------- write ----------
wb = Workbook()
ws = wb.active
ws.title = "在途可售SKU汇总"
headers = ["SKU", "SPU(产品表)", "品名(产品表)",
           "在途-仓库A(5775)", "可售-仓库A(5775)",
           "在途-仓库B(WPLA)", "可售-仓库B(WPLA)",
           "在途-仓库C(休斯顿)", "可售-仓库C(休斯顿)",
           "在途合计", "可售合计"]
ws.append(headers)
for _, r in kept.iterrows():
    ws.append([r["SKU"], r["SPU(产品表)"], r["品名(产品表)"],
               r["在途-仓库A(5775)"], r["可售-仓库A(5775)"],
               r["在途-仓库B(WPLA)"], r["可售-仓库B(WPLA)"],
               r["在途-仓库C(休斯顿)"], r["可售-仓库C(休斯顿)"],
               r["在途合计"], r["可售合计"]])

tot = {c: int(kept[c].sum()) for c in wh_cols}
ws.append(["合计", "", f"{len(kept)}个SKU",
           tot["在途-仓库A(5775)"], tot["可售-仓库A(5775)"],
           tot["在途-仓库B(WPLA)"], tot["可售-仓库B(WPLA)"],
           tot["在途-仓库C(休斯顿)"], tot["可售-仓库C(休斯顿)"],
           tot["在途合计"], tot["可售合计"]])

# style
hdr_fill = PatternFill("solid", fgColor="4472C4")
hdr_font = Font(bold=True, color="FFFFFF", name="微软雅黑", size=10)
tot_fill = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")
left = Alignment(horizontal="left", vertical="center")
for c in range(1, len(headers)+1):
    cell = ws.cell(row=1, column=c)
    cell.fill = hdr_fill; cell.font = hdr_font; cell.alignment = center; cell.border = border
last = ws.max_row
for c in range(1, len(headers)+1):
    cell = ws.cell(row=last, column=c)
    cell.fill = tot_fill; cell.font = Font(bold=True, name="微软雅黑"); cell.border = border
    cell.alignment = center if c >= 4 else left
for rr in range(2, ws.max_row+1):
    for c in range(4, len(headers)+1):
        ws.cell(row=rr, column=c).number_format = "#,##0"
        ws.cell(row=rr, column=c).alignment = center
    for c in (1, 2, 3):
        ws.cell(row=rr, column=c).alignment = left
widths = [32, 16, 26, 14, 14, 14, 14, 16, 16, 10, 10]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"

# ---------- 说明 sheet (replace) ----------
ws2 = wb.create_sheet("说明")
matched = int((kept["SPU(产品表)"].astype(str).str.strip() != "").sum())
notes = [
    "在途/可售 SKU 汇总说明（已删除在途&可售均为0的空记录）",
    "",
    f"1. 数据源：仓库A=5775 / 仓库B=WPLA / 仓库C=TX休斯顿大牛仓。",
    f"2. 原始 SKU {before} 个，已删除「在途合计=0 且 可售合计=0」的行 {removed} 个，保留 {len(kept)} 个。",
    f"3. 保留的 SKU 中 {matched} 个匹配到产品表 SPU/品名。",
    f"4. 在途合计 {tot['在途合计']:,} ，可售合计 {tot['可售合计']:,} 。",
    f"   分仓：仓库A 在途 {tot['在途-仓库A(5775)']:,}/可售 {tot['可售-仓库A(5775)']:,}；仓库B 在途 {tot['在途-仓库B(WPLA)']:,}/可售 {tot['可售-仓库B(WPLA)']:,}；仓库C 在途 {tot['在途-仓库C(休斯顿)']:,}/可售 {tot['可售-仓库C(休斯顿)']:,}。",
    "5. 仓库C(休斯顿)导出仅含「在途」一列，无「可售」字段，故仓库C 可售记为 0。",
    "6. 口径：在途=采购/调拨未入库；可售=在库可销售。排序：按 (在途+可售) 降序。",
]
for i, t in enumerate(notes, 1):
    ws2.cell(row=i, column=1, value=t)
ws2.column_dimensions["A"].width = 115

wb.save(OUT)
print("SAVED:", OUT)
print("before:", before, "removed:", removed, "kept:", len(kept))
print("在途合计:", tot["在途合计"], "可售合计:", tot["可售合计"])
