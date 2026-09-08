import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = "/Users/temu/WorkBuddy/2026-08-31-20-57-42/在途可售SKU汇总.xlsx"

df = pd.read_excel(SRC, sheet_name="在途可售SKU汇总", dtype=str)
df = df[df["SKU"] != "合计"].copy()  # drop totals row

cols = ["在途-仓库A(5775)","可售-仓库A(5775)","在途-仓库B(WPLA)","可售-仓库B(WPLA)",
        "在途-仓库C(休斯顿)","可售-仓库C(休斯顿)","在途合计","可售合计"]
for c in cols:
    df[c] = pd.to_numeric(df[c].str.replace(",",""), errors="coerce").fillna(0).astype(int)

g = df.groupby("SPU(产品表)", as_index=False)
agg = g.agg(
    SKU数=("SKU", "count"),
    品名=("品名(产品表)", "first"),
    **{c: (c, "sum") for c in cols},
)
agg["总量"] = agg["在途合计"] + agg["可售合计"]
agg = agg.sort_values("总量", ascending=False).drop(columns="总量")

# ---------- write ----------
wb = Workbook()
ws = wb.active
ws.title = "SPU汇总"
headers = ["SPU(产品表)", "品名(产品表)", "SKU数",
           "在途-仓库A(5775)", "可售-仓库A(5775)",
           "在途-仓库B(WPLA)", "可售-仓库B(WPLA)",
           "在途-仓库C(休斯顿)", "可售-仓库C(休斯顿)",
           "在途合计", "可售合计"]
ws.append(headers)
for _, r in agg.iterrows():
    ws.append([r["SPU(产品表)"], r["品名"], int(r["SKU数"]),
               r["在途-仓库A(5775)"], r["可售-仓库A(5775)"],
               r["在途-仓库B(WPLA)"], r["可售-仓库B(WPLA)"],
               r["在途-仓库C(休斯顿)"], r["可售-仓库C(休斯顿)"],
               r["在途合计"], r["可售合计"]])

tot = {c: int(agg[c].sum()) for c in cols}
ws.append(["合计", "", int(agg["SKU数"].sum()),
           tot["在途-仓库A(5775)"], tot["可售-仓库A(5775)"],
           tot["在途-仓库B(WPLA)"], tot["可售-仓库B(WPLA)"],
           tot["在途-仓库C(休斯顿)"], tot["可售-仓库C(休斯顿)"],
           tot["在途合计"], tot["可售合计"]])

# style
hdr_fill = PatternFill("solid", fgColor="4472C4")
hdr_font = Font(bold=True, color="FFFFFF", name="微软雅黑", size=10)
tot_fill = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")
left = Alignment(horizontal="left", vertical="center")

for c in range(1, len(headers)+1):
    cell = ws.cell(row=1, column=c)
    cell.fill = hdr_fill; cell.font = hdr_font; cell.alignment = center; cell.border = border
last = ws.max_row
for c in range(1, len(headers)+1):
    cell = ws.cell(row=last, column=c)
    cell.fill = tot_fill; cell.font = Font(bold=True, name="微软雅黑"); cell.border = border
    cell.alignment = center if c >= 3 else left
for rr in range(2, ws.max_row+1):
    for c in range(3, len(headers)+1):
        ws.cell(row=rr, column=c).number_format = "#,##0"
        ws.cell(row=rr, column=c).alignment = center
    for c in (1, 2):
        ws.cell(row=rr, column=c).alignment = left

widths = [18, 30, 8, 14, 14, 14, 14, 16, 16, 10, 10]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"

out = "/Users/temu/WorkBuddy/2026-08-31-20-57-42/在途可售SPU汇总.xlsx"
wb.save(out)
print("SAVED:", out)
print("SPU 数量:", len(agg), "SKU 数合计:", int(agg["SKU数"].sum()))
print("在途合计:", tot["在途合计"], "可售合计:", tot["可售合计"])
print("分仓 在途 A/B/C:", tot["在途-仓库A(5775)"], tot["在途-仓库B(WPLA)"], tot["在途-仓库C(休斯顿)"])
print("分仓 可售 A/B/C:", tot["可售-仓库A(5775)"], tot["可售-仓库B(WPLA)"], tot["可售-仓库C(休斯顿)"])
print("Top3 SPU:")
print(agg.head(3)[["SPU(产品表)","品名","SKU数","在途合计","可售合计"]].to_string(index=False))
